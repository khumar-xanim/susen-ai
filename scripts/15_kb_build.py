"""Build the knowledge base: AMU lecture materials (PDF/DOCX) + MedlinePlus health topics (English).
Usage: python scripts\\15_kb_build.py [--amu "C:\\susen_ai\\AMU pharmacology Mühazirə mətnləri"] [--medlineplus kb_sources\\medlineplus]"""
import argparse
import hashlib
import html
import json
import re
import xml.etree.ElementTree as ET
from collections import Counter
from pathlib import Path

import numpy as np
import pymupdf
from tqdm import tqdm

from config import ROOT
from kb import KB_CHUNKS, KB_VECTORS, embed
from leaflet_parser import chunk_section

ap = argparse.ArgumentParser()
ap.add_argument("--amu", default=str(ROOT / "AMU pharmacology Mühazirə mətnləri"))
ap.add_argument("--medlineplus", default=str(ROOT / "kb_sources" / "medlineplus"))
args = ap.parse_args()


def pdf_chunks(path, subject):
    out, scanned = [], 0
    with pymupdf.open(path) as doc:
        for pno, page in enumerate(doc, start=1):
            lines = [x.strip() for x in page.get_text().splitlines() if x.strip()]
            text = " ".join(lines)
            if len(text) < 40:                       # title / "thank you" slides, or an image-only page
                scanned += bool(page.get_images())
                continue
            for ch in chunk_section({"heading": lines[0][:100], "paras": [(pno, pno, text)]}):
                out.append({"source": "amu", "title": f"{subject} / {path.stem}", "page": ch["page_start"], "text": ch["text"]})
    return out, scanned


def docx_chunks(path, subject):
    from docx import Document
    doc = Document(path)
    sections, cur = [], {"heading": path.stem, "paras": []}
    for p in doc.paragraphs:
        t = " ".join(p.text.split())
        if not t:
            continue
        if p.style is not None and p.style.name.lower().startswith("heading"):
            if cur["paras"]:
                sections.append(cur)
            cur = {"heading": t[:100], "paras": []}
        else:
            cur["paras"].append((1, 1, t))
    for tb in doc.tables:
        for r in tb.rows:
            cells = [" ".join(c.text.split()) for c in r.cells if c.text.strip()]
            if cells:
                cur["paras"].append((1, 1, " | ".join(dict.fromkeys(cells))))
    if cur["paras"]:
        sections.append(cur)
    out = []
    for s in sections:
        for ch in chunk_section(s):
            out.append({"source": "amu", "title": f"{subject} / {path.stem}", "page": 0, "text": ch["text"]})
    return out


def medlineplus_chunks(folder):
    out = []
    for path in sorted(Path(folder).glob("*.xml")):
        for el in ET.parse(path).getroot().iter():
            if el.tag.split("}")[-1] != "health-topic" or el.get("language") != "English":
                continue
            summ = next((c for c in el if c.tag.split("}")[-1] == "full-summary"), None)
            if summ is None or not (summ.text or "").strip():
                continue
            text = html.unescape(re.sub(r"<[^>]+>", " ", summ.text))
            text = " ".join(text.split())
            also = [" ".join((c.text or "").split()) for c in el if c.tag.split("}")[-1] == "also-called"]
            head = el.get("title", "") + (f" (also called: {', '.join(a for a in also if a)})" if any(also) else "")
            for ch in chunk_section({"heading": head, "paras": [(0, 0, f"{head}. {text}")]}):
                out.append({"source": "medlineplus", "title": el.get("title", ""), "url": el.get("url", ""),
                            "page": 0, "text": ch["text"]})
    return out


chunks, scanned_files, skipped = [], [], []
amu = Path(args.amu)
if amu.exists():
    for p in tqdm(sorted(amu.rglob("*")), desc="AMU files"):
        if not p.is_file():
            continue
        subject = p.parent.name
        try:
            if p.suffix.lower() == ".pdf":
                c, sc = pdf_chunks(p, subject)
                chunks += c
                if not c and sc:
                    scanned_files.append(p.name)
            elif p.suffix.lower() == ".docx":
                chunks += docx_chunks(p, subject)
            else:
                skipped.append(p.name)
        except Exception as e:
            skipped.append(f"{p.name} ({str(e)[:60]})")
else:
    print("AMU folder not found:", amu)
if Path(args.medlineplus).exists():
    chunks += medlineplus_chunks(args.medlineplus)

seen, uniq = set(), []
for c in chunks:                                     # drop exact duplicates
    h = hashlib.md5(" ".join(c["text"].split()).lower().encode()).hexdigest()
    if h not in seen:
        seen.add(h)
        uniq.append(c)
print(f"\nchunks: {len(uniq)}  by source: {dict(Counter(c['source'] for c in uniq))}")
print("by subject (AMU):", dict(Counter(c["title"].split(' / ')[0] for c in uniq if c["source"] == 'amu')))
print(f"image-only PDFs (need OCR later): {len(scanned_files)} {scanned_files[:8]}")
print(f"skipped files: {len(skipped)} {skipped[:8]}")

def cyr_ratio(t):
    letters = [c for c in t if c.isalpha()]
    return sum("\u0400" <= c <= "\u04ff" for c in letters) / max(len(letters), 1)


before = len(uniq)
uniq = [c for c in uniq if cyr_ratio(c["text"]) < 0.3]     # old Cyrillic-script texts cannot be read by the model
print(f"dropped {before - len(uniq)} chunks written in Cyrillic script")


if not uniq:
    raise SystemExit("Nothing to index.")

vecs = []
for i in tqdm(range(0, len(uniq), 64), desc="embedding"):
    vecs.append(embed([c["text"] for c in uniq[i:i + 64]]))
np.save(KB_VECTORS, np.vstack(vecs))
with open(KB_CHUNKS, "w", encoding="utf-8") as f:
    for c in uniq:
        f.write(json.dumps(c, ensure_ascii=False) + "\n")
print("saved:", KB_CHUNKS, KB_VECTORS)